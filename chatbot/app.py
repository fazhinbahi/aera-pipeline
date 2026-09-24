"""
Demand Planning Chatbot — Streamlit UI.

Run with:
    cd /Users/faizanriaz/Documents/aera/chatbot
    streamlit run app.py
"""

import io
import os
import re
import sys
import time
from datetime import datetime

import pandas as pd
import streamlit as st

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from agent import run_agent
from prework_queries import fetch_customer_analysis, fetch_accuracy, fetch_customers, _client, GCP_PROJECT, DATASET
from prework_pdf import build_prework_pdf

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Demand Planning Assistant",
    page_icon="📊",
    layout="wide",
)

_SUB_SEGMENTS = [
    "EMEA ENTERP", "EMEA DEVELOP", "EMEA GTR", "EMEA IMC",
    "APAC ENTERP", "APAC DEVELOP", "APAC GTR", "APAC IMC",
]



@st.cache_data(ttl=3600, show_spinner=False)
def _load_market_map() -> dict[str, list[str]]:
    """Returns {country: sorted list of sub-segments} from BQ (cached 1 hr)."""
    try:
        df = _client().query(
            f"SELECT DISTINCT Country_Name, Sub_Segments "
            f"FROM `{GCP_PROJECT}.{DATASET}.customer_analysis` "
            "WHERE Country_Name IS NOT NULL AND Sub_Segments IS NOT NULL "
            "ORDER BY Country_Name, Sub_Segments"
        ).to_dataframe()
        result: dict[str, list[str]] = {}
        for _, row in df.iterrows():
            result.setdefault(row["Country_Name"], []).append(row["Sub_Segments"])
        return result
    except Exception:
        return {}

# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    st.title("📊 Demand Planning Assistant")
    st.caption("Demand Planning Intelligence")
    st.divider()
    st.markdown(
        "**Available tables**\n"
        "- `customer_analysis` — order history, AdjFC, Budget, SO\n"
        "- `stat_3pd_forecast` — SF / 3PD / Source FC\n"
        "- `lag1_data` — Lag-1 / n-3 (PBI Lag-3) FA/FB vs actuals (2026)\n"
    )
    st.divider()
    st.markdown(
        "**Example questions**\n"
        "- *Prepare analysis for Australia market*\n"
        "- *Give me China Lag-3 FA/FB for August*\n"
        "- *Compare SF vs 3PD for APAC IMC in Aug 2026*\n"
        "- *Top 10 customers by volume in UK 2025*\n"
        "- *What is the YTD growth for EMEA Enterprise?*\n"
    )
    st.divider()

    # ── Pre-Work PDF Generator ────────────────────────────────────────────────
    st.subheader("📋 Generate Pre-Work PDF")
    _market_map = _load_market_map()
    _all_countries = sorted(_market_map.keys())

    pw_country = st.selectbox(
        "Country",
        options=_all_countries,
        index=None,
        placeholder="Select a country…",
        key="pw_country",
    )

    # Sub-segment list trims to only what exists for the selected country
    _available_subsegs = _market_map.get(pw_country, _SUB_SEGMENTS) if pw_country else _SUB_SEGMENTS
    pw_subseg = st.selectbox(
        "Sub-Segment",
        options=_available_subsegs,
        index=None,
        placeholder="Select a sub-segment…",
        key="pw_subseg",
    )

    # Customer multi-select loads once both country + sub-segment are chosen
    pw_customer_nums  = []
    pw_customer_names = []
    if pw_country and pw_subseg:
        try:
            _customers = fetch_customers(pw_country, pw_subseg)
        except Exception as _cust_err:
            st.warning(f"Could not load customers: {_cust_err}")
            _customers = []
        _cust_options = [f"{name}  ({num})" for num, name in _customers]
        _cust_sel = st.multiselect(
            "Customer (optional)",
            options=_cust_options,
            default=[],
            placeholder="All customers…",
            key="pw_customer",
        )
        for _sel in _cust_sel:
            pw_customer_nums.append(_sel.rsplit("(", 1)[-1].rstrip(")"))
            pw_customer_names.append(_sel.rsplit("  (", 1)[0])
    else:
        st.multiselect("Customer (optional)", options=[],
                       disabled=True, key="pw_customer",
                       placeholder="Select country & sub-segment first…")

    if st.button("Generate Pre-Work", type="primary", use_container_width=True):
        if not pw_country or not pw_subseg:
            st.error("Please select both a country and a sub-segment.")
        else:
            _cust_label = (", ".join(pw_customer_names) if pw_customer_names else None)
            scope = f"{pw_country} {pw_subseg}" + (f" — {_cust_label}" if _cust_label else "")
            with st.spinner(f"Building pre-work for {scope}…"):
                try:
                    ca  = fetch_customer_analysis(pw_country, pw_subseg,
                                                  pw_customer_nums or None)
                    acc = fetch_accuracy(pw_country, pw_subseg,
                                        pw_customer_nums or None)
                    if ca.empty:
                        st.warning(
                            f"No data found for **{pw_country}** / **{pw_subseg}**. "
                            "Check the country name matches exactly (e.g. 'United Kingdom', "
                            "'Utd.Arab Emir.', 'Australia')."
                        )
                    else:
                        pdf_bytes = build_prework_pdf(ca, acc, pw_country, pw_subseg,
                                                      customer_name=_cust_label)
                        safe = lambda s: s.replace(' ', '_').replace(',', '')[:40] if s else ''
                        fname = (f"PreWork_{safe(pw_country)}_{safe(pw_subseg)}"
                                 + (f"_{safe(_cust_label)}" if _cust_label else "")
                                 + ".pdf")
                        st.download_button(
                            label="⬇ Download PDF",
                            data=pdf_bytes,
                            file_name=fname,
                            mime="application/pdf",
                            use_container_width=True,
                        )
                        st.success(f"PDF ready — {len(ca):,} rows processed.")
                except Exception as exc:
                    st.error(f"Error generating PDF: {exc}")

    st.divider()
    if st.button("🗑 Clear conversation"):
        st.session_state.messages = []
        st.session_state.chat_display = []
        st.session_state.pop("_xls_cache", None)
        st.rerun()

# ── Session state init ────────────────────────────────────────────────────────
if "messages" not in st.session_state:
    st.session_state.messages: list = []       # Anthropic API message history
if "chat_display" not in st.session_state:
    st.session_state.chat_display: list = []   # UI display history


def _sheet_name(title, idx: int, used: set) -> str:
    """Excel-safe, unique sheet name (Excel caps these at 31 chars and bans []:*?/\\)."""
    base = re.sub(r"[\[\]:*?/\\]", "-", str(title or "")).strip()
    base = re.sub(r"\s+", " ", base)[:31] or f"Table {idx + 1}"
    name, n = base, 1
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def _excel_bytes(dataframes: list) -> bytes:
    """Every result table of one answer in a single workbook, one tab per table."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf, used = io.BytesIO(), set()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for i, dfi in enumerate(dataframes):
            df = dfi["df"]
            sheet = _sheet_name(dfi.get("title"), i, used)
            df.to_excel(xw, sheet_name=sheet, index=False)
            ws = xw.sheets[sheet]
            for c in range(len(df.columns)):
                col_name = str(df.columns[c])
                cell = ws.cell(row=1, column=c + 1)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1B2B4B")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                # Width from the header and a sample of values. Positional access
                # (duplicate column labels would otherwise return a DataFrame) and
                # str() per value — BigQuery's nullable dtypes keep NA after
                # astype(str), and len(NA) raises TypeError.
                try:
                    longest = max((len(str(v)) for v in df.iloc[:200, c].tolist()),
                                  default=0)
                except Exception:
                    longest = 0
                ws.column_dimensions[get_column_letter(c + 1)].width = \
                    min(max(longest, len(col_name)) + 2, 42)
            ws.freeze_panes = "A2"
    return buf.getvalue()


def _excel_for(dataframes: list, chat_idx: int) -> bytes:
    """Memoised per message — Streamlit reruns the whole script on every interaction."""
    cache = st.session_state.setdefault("_xls_cache", {})
    if chat_idx not in cache:
        cache[chat_idx] = _excel_bytes(dataframes)
    return cache[chat_idx]


def _render_results(dataframes: list, chat_idx: int):
    """Tables, a CSV per table, and one multi-tab Excel for the whole answer."""
    if not dataframes:
        return
    for dfi in dataframes:
        title = dfi.get("title", "")
        if title:
            st.caption(title)
        st.dataframe(dfi["df"], use_container_width=True, hide_index=True)

    # The answer itself is the deliverable — a download problem must never take
    # the page down with it.
    try:
        xls = _excel_for(dataframes, chat_idx)
    except Exception as exc:
        st.caption(f"⚠ Excel export unavailable for this answer ({type(exc).__name__}). "
                   "The tables above are still complete.")
        return

    n = len(dataframes)
    st.download_button(
        label=f"⬇ Download Excel ({n} tab{'s' if n > 1 else ''})",
        data=xls,
        file_name=f"demand_planning_{datetime.now():%Y-%m-%d_%H%M}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=f"xls_{chat_idx}",
    )


# ── Welcome message (shown only before first user message) ───────────────────
if not st.session_state.chat_display:
    with st.chat_message("assistant"):
        st.markdown(
            "## Hi, I'm the Aera Demand Planning Bot 👋\n\n"
            "I'm connected to your live demand planning data and can help you analyse "
            "forecast performance, customer trends, and market insights across "
            "**EMEA** and **APAC** in real time.\n\n"
            "**Here's what I can do:**\n\n"
            "- 📦 **Volume & actuals** — YTD sales, YoY growth, top customers by market\n"
            "- 🎯 **Forecast accuracy** — Lag-1 / n-3 WMAPE & Bias vs actuals (2026, PBI convention)\n"
            "- 📐 **Plan alignment** — AdjFC vs Budget vs PMCF, confirmed SO vs plan\n"
            "- 📊 **Forecast comparison** — AdjFC vs SF vs 3PD vs Source Forecast\n"
            "- 🚨 **Deviation flags** — which sub-brands are over/under plan and by how much\n"
            "- 📋 **Pre-work PDF** — generate a full pre-alignment document for any market "
            "(use the sidebar →)\n\n"
            "**Try asking:**\n"
            "> *What is the YTD volume for Australia APAC IMC?*\n\n"
            "> *Where do SF and consensus disagree most in Japan?*\n\n"
            "> *Compare AdjFC vs SO for UK in H2 2026*\n\n"
            "> *Top 10 sub-brands by 2026 actual sales in EMEA Enterprise*"
        )

# ── Render existing conversation ──────────────────────────────────────────────
for chat_idx, item in enumerate(st.session_state.chat_display):
    with st.chat_message(item["role"]):
        st.markdown(item["content"])

        _render_results(item.get("dataframes", []), chat_idx)


# ── Chat input ────────────────────────────────────────────────────────────────
prompt = st.chat_input("Ask anything about your demand planning data…")

if prompt:
    # Show user message immediately
    with st.chat_message("user"):
        st.markdown(prompt)
    st.session_state.chat_display.append({"role": "user", "content": prompt, "dataframes": []})

    # Add to agent history
    st.session_state.messages.append({"role": "user", "content": prompt})

    # Run agent, narrating each step so the wait is visible rather than blank
    with st.chat_message("assistant"):
        started = time.time()
        with st.status("Working on it…", expanded=True) as status:

            def _on_event(kind: str, d: dict):
                el = f"{time.time() - started:.0f}s"
                if kind == "thinking":
                    status.update(label=f"Thinking… ({el})")
                    if d.get("step", 1) > 1:
                        st.write(f"🤔 Reviewing what came back… `{el}`")
                    else:
                        st.write("🤔 Working out how to answer this…")
                elif kind == "plan":
                    txt = " ".join(d.get("text", "").split())
                    if txt:
                        st.write(f"💭 {txt[:300]}{'…' if len(txt) > 300 else ''}")
                elif kind == "schema":
                    st.write(f"📂 Checking the columns in `{d.get('table','')}`")
                elif kind == "sql":
                    status.update(label=f"Querying BigQuery… ({el})")
                    st.write(f"🔎 Running query — *{d.get('label','')}*")
                elif kind == "sql_done":
                    if d.get("error"):
                        st.write(f"　↳ ⚠ query failed, adjusting: "
                                 f"{str(d['error'])[:120]}")
                    else:
                        st.write(f"　↳ ✅ {d.get('rows', 0):,} rows back `{el}`")
                elif kind == "writing":
                    status.update(label=f"Writing up the answer… ({el})")
                    st.write("📝 Pulling it together…")

            try:
                result = run_agent(st.session_state.messages, on_event=_on_event)
                text = result["text"]
                dfs  = result["dataframes"]
                status.update(label=f"Done in {time.time() - started:.0f}s",
                              state="complete", expanded=False)
            except Exception as exc:
                text = f"⚠ Error: {exc}"
                dfs  = []
                status.update(label="Something went wrong", state="error",
                              expanded=True)

        st.markdown(text)

        chat_idx = len(st.session_state.chat_display)
        _render_results(dfs, chat_idx)

    # Persist to display history
    st.session_state.chat_display.append({
        "role": "assistant",
        "content": text,
        "dataframes": dfs,
    })
